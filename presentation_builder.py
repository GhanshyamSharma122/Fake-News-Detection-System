from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUTPUT = "NewsLens_AAT_Presentation.pptx"
W, H = 13.333, 7.5
BG = RGBColor(16, 20, 26)
PANEL = RGBColor(28, 35, 44)
WHITE = RGBColor(238, 240, 243)
MUTED = RGBColor(174, 184, 196)
GREEN = RGBColor(198, 244, 50)
ORANGE = RGBColor(242, 169, 0)
RED = RGBColor(229, 56, 59)


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, x, y, w, h, text, size=14, color=WHITE, font="Aptos", bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, x, y, w, h, lines, size=12, color=WHITE, font="Aptos", line_spacing=1.15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(8)
        paragraph.line_spacing = line_spacing
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_panel(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(0.7)
    return shape


def add_line(slide, x1, y1, x2, y2, color=MUTED, width=1):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_header(slide, number, title, page):
    badge = add_panel(slide, 0.6, 0.3, 0.85, 0.48, GREEN, GREEN)
    add_text(slide, 0.6, 0.3, 0.85, 0.48, number, 12, BG, "Courier New", True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, 1.75, 0.31, 10.4, 0.5, title, 25, WHITE, "Cambria", True)
    add_text(slide, 11.95, 7.03, 0.65, 0.22, page, 9, MUTED, "Courier New", False, PP_ALIGN.RIGHT)


def add_label(slide, x, y, text, color=GREEN):
    add_text(slide, x, y, 5.6, 0.24, text.upper(), 9, color, "Courier New", True)


def add_arrow(slide, x, y, w=0.38, h=0.22, color=GREEN):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    # Cover layout follows the reference presentation: centered title,
    # presenter row, guide row, and institutional footer. The background stays
    # consistent with the other NewsLens slides.
    add_text(slide, 1.0, 0.72, 11.35, 1.75,
             "NewsLens: Explainable Fake News\nDetection Using Natural Language\nProcessing",
             27, WHITE, "Cambria", True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.01)
    add_line(slide, 3.9, 2.78, 6.1, 2.78, GREEN, 1.4)
    badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.HEXAGON, Inches(6.25), Inches(2.56), Inches(0.55), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BG
    badge.line.color.rgb = GREEN
    badge.line.width = Pt(1.2)
    add_text(slide, 6.25, 2.59, 0.55, 0.33, "N", 12, GREEN, "Courier New", True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.01)
    add_line(slide, 7.05, 2.78, 9.43, 2.78, GREEN, 1.4)
    add_text(slide, 0.8, 3.12, 11.7, 0.25, "PRESENTED BY", 9, GREEN, "Courier New", True, PP_ALIGN.CENTER)
    presenters = [
        ("Hitesh Sharma", "1BM23CS114", 0.65),
        ("Ghanshyam Sharma", "1BM23CS100", 3.45),
        ("Shreyash Shaurya", "1BM23CS354", 6.25),
        ("Utkrisht Umang", "1BM23CS355", 9.05),
    ]
    for name, usn, x in presenters:
        add_text(slide, x, 3.55, 2.35, 0.3, name, 11, WHITE, "Aptos", True, PP_ALIGN.CENTER)
        add_text(slide, x, 3.94, 2.35, 0.23, usn, 8.5, MUTED, "Courier New", False, PP_ALIGN.CENTER)
    add_text(slide, 0.8, 4.55, 11.7, 0.23, "UNDER THE GUIDANCE OF", 8, GREEN, "Courier New", True, PP_ALIGN.CENTER)
    add_text(slide, 0.8, 4.82, 11.7, 0.3, "Rajeshwari K", 15, WHITE, "Aptos", True, PP_ALIGN.CENTER)
    add_text(slide, 0.8, 5.18, 11.7, 0.23, "Associate Professor", 10, MUTED, "Aptos", False, PP_ALIGN.CENTER)
    add_text(slide, 0.8, 5.72, 11.7, 0.3, "Department of Computer Science and Engineering", 12, WHITE, "Aptos", False, PP_ALIGN.CENTER)
    add_text(slide, 0.8, 6.14, 11.7, 0.3, "B.M.S. College of Engineering", 11, MUTED, "Aptos", False, PP_ALIGN.CENTER)


def slide_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "01", "Problem statement and objectives", "2")
    add_label(slide, 0.6, 1.58, "PROBLEM")
    add_text(slide, 0.6, 1.98, 5.4, 1.25, "A user may see a headline before checking the source. NewsLens gives a quick language-based signal for deciding whether the text needs verification.", 22, WHITE, "Cambria", True)
    add_text(slide, 0.6, 3.48, 5.2, 1.25, "The output is not a fact-check. It is a prediction learned from labelled examples, so a person must still check the publisher, date, author and supporting evidence.", 13, MUTED)
    add_label(slide, 6.95, 1.58, "OBJECTIVES")
    objectives = [
        "Train with Fake.csv and True.csv and evaluate on a held-out split.",
        "Use one model for headlines and another for longer articles.",
        "Return a score, model name, language signals and a caution message.",
    ]
    for i, objective in enumerate(objectives, 1):
        yy = 2.05 + (i - 1) * 1.22
        add_line(slide, 6.95, yy - 0.12, 12.05, yy - 0.12, MUTED, 0.7)
        add_text(slide, 6.95, yy, 0.62, 0.35, f"0{i}", 16, GREEN, "Courier New", True)
        add_text(slide, 7.7, yy - 0.02, 4.35, 0.68, objective, 13, WHITE, "Aptos", False)
    add_panel(slide, 0.6, 5.75, 11.45, 0.78, PANEL, PANEL)
    add_text(slide, 0.9, 5.98, 10.9, 0.28, "Scope: English text · two labels · local model artifact · Flask API · browser interface", 12, GREEN, "Courier New", True)


def slide_related(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "02", "NLP terms used in NewsLens", "3")
    add_label(slide, 0.6, 1.58, "KEY VOCABULARY", ORANGE)
    related = [
        ("Tokenization", "Splitting the input into words so the vectorizer can process the text."),
        ("TF-IDF", "Giving more weight to useful words and less weight to words common in many documents."),
        ("Classification", "Learning a relationship between text features and the two dataset labels."),
    ]
    for i, (title, body) in enumerate(related):
        yy = 2.02 + i * 1.1
        add_line(slide, 0.6, yy - 0.12, 5.55, yy - 0.12, MUTED, 0.7)
        add_text(slide, 0.6, yy, 5.7, 0.28, title, 14, WHITE, "Cambria", True)
        add_text(slide, 0.6, yy + 0.34, 5.5, 0.54, body, 10, MUTED)
    add_panel(slide, 6.55, 1.55, 5.5, 4.95, PANEL, PANEL)
    add_text(slide, 6.95, 1.95, 4.6, 0.25, "IMPLEMENTATION CHOICES", 10, GREEN, "Courier New", True)
    add_text(slide, 6.95, 2.45, 4.5, 0.35, "Why this baseline?", 17, WHITE, "Cambria", True)
    add_text(slide, 6.95, 2.86, 4.55, 0.85, "TF-IDF with Logistic Regression is fast to train, simple to inspect and practical for a college web application.", 12, MUTED)
    add_line(slide, 6.95, 4.05, 11.65, 4.05, MUTED, 0.7)
    add_text(slide, 6.95, 4.38, 4.5, 0.35, "Two input modes", 17, WHITE, "Cambria", True)
    add_text(slide, 6.95, 4.8, 4.55, 1.05, "A headline uses title text only. A longer input uses title plus article text. The API shows which model was selected.", 12, MUTED)
    add_text(slide, 0.6, 6.75, 11.2, 0.28, "Important: a high score means the language resembles the fake class in this dataset; it does not prove that the claim is false.", 10, GREEN, "Courier New", True)


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "03", "System architecture and model flow", "4")
    add_label(slide, 0.6, 1.35, "TRAINING PIPELINE")
    steps = [("DATA", "Fake.csv\nTrue.csv"), ("NLP", "tokenize\ntext"), ("FEATURES", "word\nTF-IDF"), ("MODELS", "headline +\narticle"), ("ARTIFACT", "joblib\nmetadata")]
    start_x = 0.65
    for i, (title, body) in enumerate(steps):
        x = start_x + i * 2.4
        add_panel(slide, x, 1.75, 1.85, 1.05, PANEL, GREEN)
        add_text(slide, x + 0.1, 1.9, 1.65, 0.2, title, 8, GREEN, "Courier New", True, PP_ALIGN.CENTER)
        add_text(slide, x + 0.1, 2.18, 1.65, 0.43, body, 11, WHITE, "Aptos", True, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_arrow(slide, x + 1.95, 2.17, 0.32, 0.2)
    add_label(slide, 0.6, 3.35, "INFERENCE PIPELINE", ORANGE)
    inference = [("USER TEXT", "headline\nor article"), ("VALIDATE", "type +\nlength"), ("ROUTE", "≤60 words\nheadline"), ("PREDICT", "score +\nsignals"), ("VERIFY", "check\nsource")]
    for i, (title, body) in enumerate(inference):
        x = start_x + i * 2.4
        fill = RGBColor(45, 39, 27) if i == 2 else PANEL
        line = ORANGE if i == 2 else MUTED
        add_panel(slide, x, 3.75, 1.85, 1.05, fill, line)
        add_text(slide, x + 0.1, 3.9, 1.65, 0.2, title, 8, ORANGE if i == 2 else GREEN, "Courier New", True, PP_ALIGN.CENTER)
        add_text(slide, x + 0.1, 4.18, 1.65, 0.43, body, 11, WHITE, "Aptos", True, PP_ALIGN.CENTER)
        if i < len(inference) - 1:
            add_arrow(slide, x + 1.95, 4.17, 0.32, 0.2, ORANGE if i == 2 else GREEN)
    add_panel(slide, 0.65, 5.45, 11.45, 0.78, PANEL, PANEL)
    add_text(slide, 0.95, 5.68, 10.8, 0.28, "Run: python training/train_model.py --input data  ·  Start: python server.py  ·  API: POST /api/analyze", 10, WHITE, "Courier New", True, PP_ALIGN.CENTER)


def slide_results(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, "04", "Evaluation and conclusions", "5")
    add_label(slide, 0.6, 1.45, "OFFLINE EVALUATION", ORANGE)
    cards = [("ARTICLE MODEL", "99.31%", "accuracy", "F1: 99.34% · held-out split", GREEN), ("HEADLINE MODEL", "96.16%", "accuracy", "F1: 96.29% · held-out split", ORANGE), ("REAL-WORLD USE", "UNKNOWN", "guaranteed accuracy", "New samples can be misclassified", MUTED)]
    for i, (label, value, small, detail, col) in enumerate(cards):
        x = 0.6 + i * 3.92
        add_panel(slide, x, 1.85, 3.52, 1.35, PANEL, PANEL)
        add_text(slide, x + 0.25, 2.08, 3.0, 0.2, label, 8, col, "Courier New", True)
        add_text(slide, x + 0.25, 2.36, 1.75, 0.38, value, 25, WHITE, "Cambria", True)
        add_text(slide, x + 2.0, 2.45, 1.15, 0.22, small, 9, MUTED, "Courier New")
        add_text(slide, x + 0.25, 2.92, 3.0, 0.2, detail, 9, MUTED)
    add_label(slide, 0.6, 3.65, "OBSERVATIONS", ORANGE)
    observations = [
        "The dataset contains 44,898 usable articles; 35,918 were used for training and 8,980 for testing.",
        "The percentages are results from a random split of this dataset, not a promise for unseen news.",
        "A new sample can be wrong because of topic, source, wording, satire or dataset bias.",
    ]
    for i, text in enumerate(observations, 1):
        yy = 4.0 + (i - 1) * 0.52
        add_text(slide, 0.6, yy, 0.4, 0.23, f"0{i}", 10, GREEN, "Courier New", True)
        add_text(slide, 1.1, yy - 0.02, 5.45, 0.37, text, 10, WHITE)
    add_panel(slide, 6.95, 3.65, 5.1, 2.35, PANEL, PANEL)
    add_text(slide, 7.3, 3.95, 4.4, 0.2, "CONCLUSION", 9, GREEN, "Courier New", True)
    add_text(slide, 7.3, 4.3, 4.3, 0.88, "These metrics describe the experiment, not the truth of a new article.", 18, WHITE, "Cambria", True)
    add_text(slide, 7.3, 5.35, 4.3, 0.38, "Use the score to start verification, not to replace it.", 11, ORANGE, "Courier New", True)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide_cover(prs)
    slide_problem(prs)
    slide_related(prs)
    slide_architecture(prs)
    slide_results(prs)
    prs.save(OUTPUT)


if __name__ == "__main__":
    build()
